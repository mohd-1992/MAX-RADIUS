"""Versioned voucher workbooks and private, offline HTML-to-PDF rendering."""
import datetime
import io
import mimetypes
import os
import threading
from decimal import Decimal, InvalidOperation
from pathlib import Path
from urllib.parse import unquote, urlsplit

FORMAT = 'MAX-VOUCHERS-1'
SNAPSHOTS = (
    'snap_price', 'snap_cost', 'snap_volume_quota_mb', 'snap_uptime_limit_mins',
    'snap_validity_value', 'snap_validity_unit', 'snap_validity_days',
    'snap_rate_download', 'snap_rate_upload', 'snap_rate_limit_str',
    'snap_simultaneous_sessions', 'snap_mikrotik_group',
)
FIELDS = ('export_format', 'username', 'password', 'voucher_pin', 'package',
          'serial_number', 'status', 'first_used_at', 'last_renewed_at', 'expires_at',
          'extra_quota_mb', 'mac_binding', 'pause_reason', 'policy_frozen_at', *SNAPSHOTS,
          'download_bytes', 'upload_bytes', 'uptime_seconds', 'batch_name')
_PDF_SLOT = threading.BoundedSemaphore(1)


def batch_download_name(name, extension):
    """Keep Unicode batch names; replace only filesystem-unsafe characters."""
    import re
    value = re.sub(r'[<>:"/\\|?*\x00-\x1f\x7f]', '-', str(name or '')).strip().rstrip('. ')
    value = value[:150].rstrip('. ') or 'حزمة الكروت'
    if value.split('.', 1)[0].upper() in {'CON', 'PRN', 'AUX', 'NUL', *(f'COM{i}' for i in range(1, 10)), *(f'LPT{i}' for i in range(1, 10))}:
        value = '_' + value
    return value + '.' + extension


def export_snapshot(batch_id, owner=None):
    """Contract + cycle counters in one consistent database read transaction."""
    from database.db import get_connection, is_mysql_conn, adapt_query
    conn = get_connection()
    cur = conn.cursor()
    try:
        if is_mysql_conn(conn):
            cur.execute('SET TRANSACTION ISOLATION LEVEL REPEATABLE READ')
            cur.execute('START TRANSACTION WITH CONSISTENT SNAPSHOT, READ ONLY')
        else:
            cur.execute('BEGIN')
        def rows(sql, args=()):
            cur.execute(adapt_query(sql, conn), args)
            names = [c[0] for c in cur.description]
            return [dict(r) if isinstance(r, dict) else dict(zip(names, r)) for r in cur.fetchall()]
        batch_rows = rows('SELECT * FROM wisp_voucher_batches WHERE id=?', (batch_id,))
        if not batch_rows or (owner is not None and batch_rows[0]['reseller_id'] != owner):
            raise ValueError('الحزمة غير موجودة.')
        fallback = ','.join('p.' + k[5:] + ' AS fallback_' + k for k in SNAPSHOTS if k != 'snap_rate_limit_str')
        cards = rows('SELECT v.*,p.name AS package_name,' + fallback +
                     ' FROM wisp_vouchers v JOIN wisp_packages p ON p.id=v.package_id WHERE v.batch_id=? ORDER BY v.id LIMIT 20001', (batch_id,))
        if not cards or len(cards) > 20000:
            raise ValueError('تصدير Excel متاح للحزم من 1 إلى 20000 كرت.')
        cycle = 'COALESCE(s.last_renewed_at,s.first_used_at,s.created_at)'
        def delta(column, baseline):
            return f'CASE WHEN b.radacctid IS NULL AND ({cycle} IS NULL OR a.acctstarttime>={cycle}) THEN COALESCE(a.{column},0) WHEN b.radacctid IS NOT NULL THEN CASE WHEN COALESCE(a.{column},0)>COALESCE(b.{baseline},0) THEN COALESCE(a.{column},0)-COALESCE(b.{baseline},0) ELSE 0 END ELSE 0 END'
        totals = {}
        for offset in range(0, len(cards), 400):
            ids = [c['id'] for c in cards[offset:offset + 400]]
            sql = f'''SELECT LOWER(s.username) AS username,
                SUM({delta('acctinputoctets','baseline_input_bytes')}) AS up_bytes,
                SUM({delta('acctoutputoctets','baseline_output_bytes')}) AS down_bytes,
                SUM({delta('acctsessiontime','baseline_seconds')}) AS uptime_secs,
                MAX(CASE WHEN a.acctstarttime<{cycle} AND (a.acctstoptime IS NULL OR a.acctstoptime>{cycle})
                     AND b.radacctid IS NULL THEN 1 ELSE 0 END) AS incomplete
                FROM wisp_vouchers s LEFT JOIN radacct a ON LOWER(a.username)=LOWER(s.username)
                LEFT JOIN wisp_session_baselines b ON b.radacctid=a.radacctid AND b.renewed_at={cycle}
                WHERE s.id IN ({','.join('?' for _ in ids)}) GROUP BY s.id,s.username'''
            for record in rows(sql, tuple(ids)):
                if record['incomplete']:
                    raise ValueError('محاسبة دورة غير مكتملة؛ لم يُصدر الملف.')
                totals[record['username']] = {key: int(record[key] or 0) for key in ('up_bytes','down_bytes','uptime_secs')}
        from core.rate_limit import build_mikrotik_rate_limit
        for card in cards:
            for key in SNAPSHOTS:
                if key != 'snap_rate_limit_str' and card.get(key) is None:
                    card[key] = card.get('fallback_' + key)
            if card.get('snap_rate_limit_str') is None:
                card['snap_rate_limit_str'] = build_mikrotik_rate_limit(download=card['snap_rate_download'],upload=card['snap_rate_upload']) or '0/0'
        return batch_rows[0], cards, totals
    finally:
        conn.rollback()
        cur.close()
        conn.close()


def make_workbook(batch, cards, totals):
    """All credentials are literal text (including =, + and leading zeros)."""
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Cards'
    ws.sheet_view.rightToLeft = True
    ws.append(FIELDS)
    for card in cards:
        accounting = totals.get(card['username'].lower())
        if accounting is None:
            raise ValueError('تعذر تصدير استهلاك إحدى الكروت؛ راجع محاسبة الدورة أولًا.')
        row = dict(card, export_format=FORMAT, voucher_pin=card.get('pin_code') or card['username'],
                   package=card['package_name'], mac_binding=card.get('bound_mac') or '',
                   download_bytes=accounting['down_bytes'], upload_bytes=accounting['up_bytes'],
                   uptime_seconds=accounting['uptime_secs'], batch_name=batch['name'])
        validate_export_row({key: '' if row.get(key) is None else row[key] for key in FIELDS})
        ws.append([None] * len(FIELDS))
        for col, name in enumerate(FIELDS, 1):
            value = row.get(name)
            if isinstance(value, datetime.datetime):
                value = value.strftime('%Y-%m-%d %H:%M:%S')
            cell = ws.cell(ws.max_row, col)
            # Even large counters stay exact beyond Excel's 15-digit precision.
            cell.value = '' if value is None else str(value)
            cell.data_type = 's'
            cell.number_format = '@'
    for cell in ws[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='1E293B')
        cell.alignment = Alignment(horizontal='center')
    from openpyxl.utils import get_column_letter
    for col in range(1, len(FIELDS) + 1):
        ws.column_dimensions[get_column_letter(col)].width = 24
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions
    info = wb.create_sheet('Information')
    info.append(['format', FORMAT])
    info.append(['batch_name', batch['name']])
    info.append(['card_count', len(cards)])
    info.append(['notice', 'يحوي الملف بيانات دخول حساسة. احتفظ به في مكان آمن.'])
    # Never treat a user supplied batch name as an Excel formula.
    for row in info:
        for cell in row:
            if isinstance(cell.value, str):
                cell.data_type = 's'
    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    return out


def validate_export_row(row):
    """Strict restore format; no guessing dates, units or missing counters."""
    if row.get('export_format') != FORMAT or any(key not in row for key in FIELDS):
        raise ValueError('صيغة ملف الكروت أو أعمدته غير مكتملة.')
    from services.import_state_service import parse_import_datetime
    result = dict(row)
    for key in ('first_used_at', 'last_renewed_at', 'expires_at', 'policy_frozen_at'):
        result[key] = parse_import_datetime(row.get(key))
    for key in (*SNAPSHOTS, 'extra_quota_mb', 'download_bytes', 'upload_bytes', 'uptime_seconds'):
        if key in ('snap_validity_unit', 'snap_rate_download', 'snap_rate_upload',
                   'snap_rate_limit_str', 'snap_mikrotik_group'):
            continue
        try:
            value = Decimal(str(row[key]))
            if not value.is_finite() or (key != 'extra_quota_mb' and value < 0):
                raise ValueError()
            if key not in ('snap_price', 'snap_cost') and value != value.to_integral_value():
                raise ValueError()
            result[key] = str(value)
        except (InvalidOperation, ValueError):
            raise ValueError('قيمة عددية غير صالحة في ملف الكروت.') from None
    if result['snap_validity_unit'] not in ('days', 'hours', 'minutes', 'months'):
        raise ValueError('وحدة الصلاحية غير صالحة.')
    if result['status'] not in ('unused', 'active', 'expired', 'suspended', 'disabled', 'recharged'):
        raise ValueError('حالة الكرت غير صالحة.')
    first, cycle, expiry = (result[k] for k in ('first_used_at', 'last_renewed_at', 'expires_at'))
    used = any(Decimal(result[k]) > 0 for k in ('download_bytes', 'upload_bytes', 'uptime_seconds'))
    if (used or result['status'] == 'active') and not (first and cycle):
        raise ValueError('دليل بداية الاستخدام والدورة غير مكتمل.')
    if cycle and (not first or cycle < first):
        raise ValueError('حد الدورة لا يطابق تاريخ الاستخدام.')
    if expiry and first and expiry < first:
        raise ValueError('تاريخ الانتهاء يسبق الاستخدام.')
    if result['status'] == 'unused' and (first or cycle or used):
        raise ValueError('الكرت غير المستخدم يحمل محاسبة أو تاريخ استخدام.')
    if not result['voucher_pin'] or not result['password'] or not result['serial_number']:
        raise ValueError('بيانات اعتماد الكرت غير مكتملة.')
    for field, limit in (('username', 64), ('password', 64), ('voucher_pin', 32), ('serial_number', 32)):
        if len(str(result.get(field) or '')) > limit:
            raise ValueError('بيانات اعتماد الكرت تتجاوز الطول المسموح.')
    return result


def render_pdf(html, static_folder):
    """No HTTP, cookies, file:// or arbitrary file reads in the renderer."""
    if not _PDF_SLOT.acquire(blocking=False):
        raise RuntimeError('جارٍ تجهيز ملف آخر. أعد المحاولة بعد قليل.')
    try:
        from playwright.sync_api import sync_playwright
        root = Path(static_folder).resolve()
        errors = []
        with sync_playwright() as pw:
            executable = os.environ.get('MAX_PDF_CHROMIUM')
            if not executable and Path('/usr/bin/chromium').exists():
                executable = '/usr/bin/chromium'
            browser = pw.chromium.launch(headless=True, executable_path=executable)
            try:
                context = browser.new_context(service_workers='block', accept_downloads=False)
                def asset(route):
                    url = urlsplit(route.request.url)
                    path = unquote(url.path)
                    if url.scheme == 'https' and url.netloc == 'max-export.invalid' and path.startswith('/static/'):
                        target = (root / path[len('/static/'):]).resolve()
                        if target.is_relative_to(root) and target.is_file():
                            route.fulfill(body=target.read_bytes(), content_type=mimetypes.guess_type(target.name)[0] or 'application/octet-stream')
                            return
                    errors.append('asset')
                    route.abort()
                context.route('**/*', asset)
                page = context.new_page()
                page.set_default_timeout(45000)
                page.on('pageerror', lambda error: errors.append('script'))
                # Relative resources use a synthetic, entirely intercepted origin.
                html = html.replace('<head>', '<head><base href="https://max-export.invalid/">', 1)
                page.set_content(html, wait_until='networkidle', timeout=60000)
                page.evaluate('''async () => {
                    await document.fonts.ready;
                    await Promise.all(Array.from(document.images).map(img => img.complete ? null :
                        new Promise(resolve => {img.onload = resolve; img.onerror = resolve;})));
                    if (Array.from(document.images).some(img => !img.naturalWidth)) throw new Error('image');
                }''')
                if errors:
                    raise RuntimeError('تعذر تحميل أحد عناصر التصميم. راجع الصورة والتصميم ثم أعد المحاولة.')
                return page.pdf(format='A4', print_background=True, prefer_css_page_size=True)
            finally:
                browser.close()
    finally:
        _PDF_SLOT.release()
