"""Batch exports/restore on disposable fixtures, no app startup or NAS."""
import ast
import datetime
import io
import re
import types
import unittest
from unittest.mock import patch

import openpyxl
from flask import current_app, send_file
from werkzeug.datastructures import FileStorage
from voucher_page_fixture import voucher_fixture, ROOT
from test_import_quota_regression import load
from test_package_page import load_functions
from services.batch_export_service import FIELDS, SNAPSHOTS, FORMAT, make_workbook, validate_export_row, batch_download_name


class BatchExportTests(unittest.TestCase):
    def setUp(self):
        self.app, self.db, self.ns, self.stubs, self.seed = voucher_fixture()
        self.modules = patch.dict('sys.modules', self.stubs)
        self.modules.start()
        self.client = self.app.test_client()
        with self.client.session_transaction() as s:
            s['manager'] = {'id': 1, 'role_id': 1}
            s['permissions'] = ['vouchers.view', 'vouchers.print']
        self.ns.update(current_app=current_app, send_file=send_file, owner_id=lambda:None)
        load_functions('web/routes/vouchers.py', ['export_batch_excel'], self.ns)
        self.app.add_url_rule('/vouchers/batch/<int:batch_id>/excel', 'export_batch_excel', self.ns['export_batch_excel'])
        db_stub = types.ModuleType('database.db')
        db_stub.query_one = self.ns['query_one']
        db_stub.is_mysql_conn = lambda conn: False
        db_stub.adapt_query = lambda sql, conn: sql
        class Connection:
            def cursor(inner): return self.db.cursor()
            def commit(inner): self.db.commit()
            def rollback(inner): self.db.rollback()
            def close(inner): pass
        db_stub.get_connection = Connection
        self.db_patch = patch.dict('sys.modules', {'database.db': db_stub})
        self.db_patch.start()
    def tearDown(self):
        self.db_patch.stop(); self.modules.stop(); self.db.close()
    def exported_row(self, **changes):
        row = dict.fromkeys(FIELDS, '')
        row.update(export_format=FORMAT, username='0000123', password='=secret', voucher_pin='0000789',
                   package='fixture-package', serial_number='000099', status='active', batch_name='fixture',
                   first_used_at='2025-01-01 00:00:00', last_renewed_at='2026-01-01 00:00:00',
                   policy_frozen_at='2025-01-01 00:00:00', expires_at='2099-01-01 00:00:00',
                   extra_quota_mb='5', download_bytes='123456789', upload_bytes='123', uptime_seconds='61')
        for field in SNAPSHOTS: row[field] = '0'
        row.update(snap_volume_quota_mb='1000', snap_uptime_limit_mins='500', snap_validity_unit='days',
                   snap_simultaneous_sessions='2', snap_mikrotik_group='ALL-SPEED', snap_rate_limit_str='2M/4M')
        row.update(changes)
        return row
    def clean(self, row):
        self.db.execute("UPDATE wisp_packages SET name='fixture-package' WHERE id=1"); self.db.commit()
        from services.batch_restore_service import validate_workbook
        ok, errors, cleaned = validate_workbook([row], 'vouchers')
        self.assertTrue(ok, errors)
        return cleaned
    def restore(self, rows):
        from services.batch_restore_service import restore_workbook
        return restore_workbook(rows)
    def test_literal_text_roundtrip_and_precision(self):
        row = self.exported_row(download_bytes='9007199254740993')
        card = dict(row, id=1, pin_code=row['voucher_pin'], package_name=row['package'])
        book = make_workbook({'name': '=batch'}, [card], {row['username']: dict(down_bytes=9007199254740993, up_bytes=123, uptime_secs=61)})
        wb = openpyxl.load_workbook(book, data_only=False)
        self.assertTrue(all(c.data_type == 's' for c in wb.active[2] if c.value is not None))
        self.assertEqual(wb['Information']['B2'].data_type, 's')
        ns = dict(io=io, re=re, openpyxl=openpyxl, xlrd=None)
        load('services/import_service.py', {'COLUMN_ALIASES','normalize_header','_excel_header_map','parse_excel_file'}, ns)
        book.seek(0)
        raw = ns['parse_excel_file'](FileStorage(book, filename='batch.xlsx'))[0]
        self.assertEqual(raw['username'], '0000123'); self.assertEqual(raw['password'], '=secret')
        self.assertEqual(raw['voucher_pin'], '0000789')
        self.assertEqual(raw['download_bytes'], '9007199254740993')
        self.assertEqual(validate_export_row(raw)['uptime_seconds'], '61')
    def test_unknown_accounting_blocks_file(self):
        with self.assertRaises(ValueError):
            make_workbook({'name':'fixture'}, [dict(username='x')], {'x':None})
    def test_strict_format_missing_and_negative(self):
        for changes in ({'export_format':'v2'}, {'upload_bytes':'-1'}, {'uptime_seconds':'1.5'},
                        {'snap_price':'NaN'}, {'last_renewed_at':''}, {'status':'unused'}):
            with self.subTest(changes=changes), self.assertRaises(ValueError):
                validate_export_row(self.exported_row(**changes))
    def test_restoration_preserves_dates_cycle_and_usage_seconds(self):
        result = self.restore(self.clean(self.exported_row()))
        self.assertTrue(result[0], result[1])
        row = dict(self.db.execute("SELECT * FROM wisp_vouchers WHERE username='0000123'").fetchone())
        self.assertEqual(row['pin_code'], '0000789'); self.assertEqual(row['extra_quota_mb'], 5)
        self.assertEqual(row['last_renewed_at'], '2026-01-01 00:00:00')
        acct = dict(self.db.execute("SELECT * FROM radacct WHERE username='0000123'").fetchone())
        self.assertEqual(acct['acctstarttime'], row['last_renewed_at'])
        self.assertEqual(acct['acctsessiontime'], 61); self.assertEqual(acct['acctoutputoctets'],123456789)
        self.assertEqual(acct['acctterminatecause'], 'Consolidated-Historical-Import')
    def test_repeat_import_refuses_all(self):
        clean = self.clean(self.exported_row())
        self.assertTrue(self.restore(clean)[0]); self.assertFalse(self.restore(clean)[0])
        self.assertEqual(self.db.execute("SELECT COUNT(*) FROM wisp_vouchers WHERE username='0000123'").fetchone()[0],1)
    def test_retained_accounting_refuses_entire_batch(self):
        clean = self.clean(self.exported_row())
        self.seed('radacct', username='0000123', acctsessionid='real-fixture')
        before = self.db.total_changes
        self.assertFalse(self.restore(clean)[0])
        self.assertEqual(self.db.total_changes, before)
    def test_expiry_and_consumed_status_cannot_be_reactivated(self):
        for status in ('expired','suspended','disabled','recharged'):
            with self.subTest(status=status):
                clean = self.clean(self.exported_row(username='state-'+status, serial_number='serial-'+status,status=status))
                self.assertTrue(self.restore(clean)[0])
                check = self.db.execute('SELECT value FROM radcheck WHERE username=? AND attribute=?', ('state-'+status,'Auth-Type')).fetchone()
                self.assertEqual(check[0], 'Reject')
    def test_expiry_since_export_is_enforced(self):
        clean = self.clean(self.exported_row(expires_at='2025-06-01 00:00:00'))
        self.assertTrue(self.restore(clean)[0])
        self.assertEqual(self.db.execute("SELECT status FROM wisp_vouchers WHERE username='0000123'").fetchone()[0], 'expired')
    def test_current_quota_includes_extra(self):
        clean = self.clean(self.exported_row(snap_volume_quota_mb='1', extra_quota_mb='5',download_bytes=str(3*1048576)))
        self.assertTrue(self.restore(clean)[0])
        self.assertEqual(self.db.execute("SELECT status FROM wisp_vouchers WHERE username='0000123'").fetchone()[0],'active')
    def test_pdf_requires_print_permission_and_batch(self):
        with self.client.session_transaction() as s: s['permissions'] = ['vouchers.view']
        self.assertEqual(self.client.get('/vouchers/print/1?export=pdf').status_code,403)
        self.assertEqual(self.client.get('/vouchers/batch/1/excel').status_code,403)
    def test_pdf_attachment_uses_same_design_html(self):
        self.db.execute("UPDATE wisp_voucher_batches SET name=? WHERE id=1", ('حزمة تجربة 001',));self.db.commit()
        with patch('services.batch_export_service.render_pdf', return_value=b'%PDF-fixture') as render:
            result = self.client.get('/vouchers/print/1?design_id=1&export=pdf')
            self.assertEqual(result.status_code,200)
            self.assertIn('attachment', result.headers['Content-Disposition'])
            from urllib.parse import unquote
            self.assertIn('حزمة تجربة 001.pdf',unquote(result.headers['Content-Disposition']))
            self.assertIn('no-store', result.headers['Cache-Control'])
            self.assertIn('customCardsPrintContainer', render.call_args.args[0])
    def test_invalid_design_does_not_silently_export_default(self):
        self.assertEqual(self.client.get('/vouchers/print/1?design_id=999&export=pdf').status_code,400)
    def test_pdf_failure_returns_no_fake_file(self):
        with patch('services.batch_export_service.render_pdf', side_effect=RuntimeError('fixture')):
            self.assertEqual(self.client.get('/vouchers/print/1?export=pdf').status_code,503)
    def test_preview_remains_html(self):
        result=self.client.get('/vouchers/print/1?design_id=1')
        self.assertEqual(result.status_code,200); self.assertIn(b'window.print()',result.data)
    def test_excel_error_is_explicit(self):
        with patch('services.batch_export_service.export_snapshot', side_effect=ValueError('fixture')):
            self.assertEqual(self.client.get('/vouchers/batch/1/excel').status_code,409)
    def test_export_snapshot_and_endpoint_all_rows(self):
        from services.batch_export_service import export_snapshot
        self.seed('wisp_voucher_batches', id=2, name='fixture export', package_id=1, card_count=1)
        row = self.exported_row()
        values = {k:row[k] for k in SNAPSHOTS}
        values.update(id=50,batch_id=2,package_id=1,username=row['username'],password=row['password'],pin_code=row['voucher_pin'],
                      serial_number=row['serial_number'],status=row['status'], first_used_at=row['first_used_at'],last_renewed_at=row['last_renewed_at'],
                      policy_frozen_at=row['policy_frozen_at'],expires_at=row['expires_at'],extra_quota_mb=5)
        self.seed('wisp_vouchers', **values)
        # Previous-cycle traffic must not leak into the exported cycle.
        self.seed('radacct',radacctid=70,username=row['username'],acctstarttime='2025-02-01',acctstoptime='2025-02-02',acctoutputoctets=999)
        self.seed('radacct',radacctid=71,username=row['username'],acctstarttime='2026-02-01',acctstoptime='2026-02-02',acctoutputoctets=123,acctsessiontime=61)
        batch,cards,totals=export_snapshot(2)
        self.assertEqual(len(cards),1); self.assertEqual(totals[row['username']]['down_bytes'],123)
        response=self.client.get('/vouchers/batch/2/excel')
        self.assertEqual(response.status_code,200)
        wb=openpyxl.load_workbook(io.BytesIO(response.data)); self.assertEqual(wb.active.max_row,2)
        self.assertEqual(wb.active.cell(2,FIELDS.index('uptime_seconds')+1).value,'61')
        self.assertIn('fixture export.xlsx',response.headers['Content-Disposition'])
        self.db.execute("UPDATE radacct SET acctstarttime='2025-12-31',acctstoptime=NULL WHERE radacctid=71");self.db.commit()
        self.assertEqual(self.client.get('/vouchers/batch/2/excel').status_code,409)
    def test_restore_failure_late_in_batch_rolls_back_everything(self):
        self.db.execute('CREATE UNIQUE INDEX fixture_serial_unique ON wisp_vouchers(serial_number)');self.db.commit()
        rows=self.clean(self.exported_row())
        other=dict(rows[0], username='second-fixture')  # duplicate serial fails after first insert
        before=self.db.execute('SELECT COUNT(*) FROM wisp_voucher_batches').fetchone()[0]
        self.assertFalse(self.restore(rows+[other])[0])
        self.assertEqual(self.db.execute('SELECT COUNT(*) FROM wisp_voucher_batches').fetchone()[0],before)
    def test_unicode_download_name_and_unsafe_characters(self):
        self.assertEqual(batch_download_name('حزمة 001 — شبكة التجارب','pdf'),'حزمة 001 — شبكة التجارب.pdf')
        self.assertEqual(batch_download_name('A/B:01?','xlsx'),'A-B-01-.xlsx')
        self.assertEqual(batch_download_name('CON','pdf'),'_CON.pdf')
        self.assertNotIn('\n', batch_download_name('a\nb','pdf'))


if __name__ == '__main__': unittest.main()
